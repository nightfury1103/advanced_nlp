from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_COLOR_INDEX


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "midterm_report_12_books_draft.docx"
TARGET = ROOT / "midterm_report_12_books_revised.docx"


def replace_paragraph(paragraph, text: str) -> None:
    """Replace paragraph content while retaining its paragraph-level formatting."""
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def replace_cell(cell, text: str) -> None:
    """Replace a table cell value and highlight the revised value."""
    paragraph = cell.paragraphs[0]
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


shutil.copy2(SOURCE, TARGET)
document = Document(TARGET)

paragraph_updates = {
    15: (
        "Kết quả mong đợi gồm 2 phần. Thứ nhất là một bộ dữ liệu văn bản có cấu trúc "
        "theo tập sách – trang – câu, trong đó mỗi đơn vị dữ liệu đều mang theo thông tin "
        "về nguồn tạo ra và trạng thái xử lý (đã đồng thuận hay còn bất đồng), tạo điều kiện "
        "cho việc hiệu đính học thuật và tái sử dụng dữ liệu về sau. Thứ hai là một bộ thực "
        "thể được gán theo năm loại PERSON, LOCATION, ORGANIZATION, DATE và TIME. Do chưa "
        "được chuyên gia kiểm định đầy đủ, bộ ngữ liệu này được xem là dữ liệu gán nhãn bán "
        "tự động (silver/pseudo-labeled), chưa phải ground truth cuối cùng."
    ),
    33: (
        "Toàn bộ 12 tập đã hoàn tất OCR agreement voting, tách câu và NER agreement voting. "
        "Từ 36.387 câu đầu vào NER, ba mô hình nhận diện được 12.564 thực thể sau voting; "
        "sau bước lọc bảo thủ, 10.054 thực thể được đưa vào bộ release candidate."
    ),
    37: (
        "Bộ dữ liệu thực thể có tên (release candidate) của 12 tập gồm 10.054 thực thể thuộc "
        "năm loại PERSON, LOCATION, ORGANIZATION, DATE và TIME; các trường hợp bất đồng được "
        "lưu riêng để phục vụ kiểm tra thủ công."
    ),
    51: (
        "Ở bước NER, 28.039 trong tổng số 36.387 câu có ít nhất một bất đồng giữa các mô hình, "
        "tương đương khoảng 77,1%. Chỉ số này cho thấy dữ liệu vẫn cần được kiểm tra thủ công, "
        "đặc biệt đối với các thực thể ít phổ biến như ORGANIZATION và TIME."
    ),
    54: (
        "Mức bất đồng ở bước NER còn cao: 28.039/36.387 câu (khoảng 77,1%) được lưu vào tập "
        "kiểm tra bất đồng."
    ),
    57: (
        "Toàn bộ 36.387 câu sau sentence agreement voting của 12 tập được đưa vào bước NER; "
        "không lấy mẫu hoặc chọn một tập con."
    ),
    62: (
        "Kết quả trên 12 tập cho thấy 28.039 câu có bất đồng NER. Nguyên nhân quan trọng là "
        "các mô hình hiện dùng đều là mô hình tổng quát, chưa được tinh chỉnh riêng cho Hán "
        "văn lịch sử Việt Nam."
    ),
    65: (
        "Đề tài đã xây dựng quy trình tạo lập ngữ liệu có cấu trúc theo ba cấp tập sách – trang "
        "– câu cho corpus 12 tập gồm 4.432 trang. OCR agreement voting và tách câu đa mô hình "
        "đã tạo ra 36.387 câu làm đầu vào NER. NER agreement voting trên toàn bộ số câu này "
        "thu được 12.564 thực thể; sau bước lọc bảo thủ, bộ release candidate còn 10.054 thực "
        "thể. Đồng thời, 28.039 câu có bất đồng được lưu riêng để tiếp tục kiểm tra."
    ),
    67: (
        "Bộ tham chiếu để chấm CER hiện vẫn là bản tạo bán tự động nhờ OCR, chưa qua xác nhận "
        "của chuyên gia Hán Nôm; các mô hình NER hiện dùng đều là mô hình đa dụng, chưa được "
        "điều chỉnh riêng cho loại văn bản này. Bộ release candidate vì vậy cần được xem là "
        "dữ liệu silver và tiếp tục kiểm định thủ công trước khi sử dụng như dữ liệu chuẩn."
    ),
}

for index, text in paragraph_updates.items():
    replace_paragraph(document.paragraphs[index], text)

# Bảng khối lượng dữ liệu xử lý.
summary_table = document.tables[1]
replace_cell(summary_table.cell(6, 1), "12.564")
replace_cell(summary_table.cell(7, 1), "10.054")
replace_cell(summary_table.cell(8, 1), "28.039")

# Bảng phân bố thực thể trong release candidate.
entity_table = document.tables[2]
entity_counts = {
    1: "5.507",
    2: "3.310",
    3: "1.026",
    4: "189",
    5: "22",
}
for row, value in entity_counts.items():
    replace_cell(entity_table.cell(row, 1), value)

document.save(TARGET)
print(TARGET)
